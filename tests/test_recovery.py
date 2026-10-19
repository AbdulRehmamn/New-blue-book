"""Regression coverage for the supplied zero-row job and normal access failures."""
import json
import queue
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extractor import extract_company, extract_listing_companies, access_problem
from storage import Store, job_folder
from crawler import Worker, Stopped
from test_scraper import FixtureWorker, FixturePage, URL, HTML

LISTING = 'https://www.thebluebook.com/search.html?class=2200&region=3'
PROFILE = URL.removesuffix('locations-contacts/')
CARDS = f'''<header>888-720-1710</header><div class="results">
<article><h3><a href="{PROFILE}">Example Builders LLC</a></h3><p>San Diego, CA 92101</p>
<a href="tel:+16195550110">Phone</a><p>office@example.org</p></article>
<article><h3><a href="https://www.thebluebook.com/iProView/888888/second/subcontractors/">Second Builders</a></h3>
<p>Austin, TX 78701</p><p>512-555-0188</p></article></div>'''


class RecoveryTests(unittest.TestCase):
    def test_company_without_h1_from_styled_name(self):
        html = HTML.replace('<h1>Example Builders LLC</h1>', '<div class="company-name">Example Builders LLC</div>')
        company = extract_company(html, URL)
        self.assertEqual(company.company_name, 'Example Builders LLC')
        self.assertEqual(company.contacts[0].name, 'Alex Morgan')

    def test_company_without_heading_from_metadata(self):
        html = '<head><meta property="og:title" content="Example Builders LLC - Locations &amp; Key Contacts | ProView"></head>' + HTML.replace('<h1>Example Builders LLC</h1>', '')
        self.assertEqual(extract_company(html, URL).company_name, 'Example Builders LLC')

    def test_schema_organization_without_h1(self):
        data = {'@type': 'GeneralContractor', 'name': 'Sample Electric', 'telephone': '619-555-0177', 'email': 'office@sample.org', 'address': {'addressLocality': 'San Diego', 'addressRegion': 'CA'}}
        html = '<script type="application/ld+json">' + json.dumps(data) + '</script>'
        company = extract_company(html, URL)
        self.assertEqual(company.company_name, 'Sample Electric')
        self.assertEqual(company.phones[0]['phone'], '+16195550177')
        self.assertEqual(company.city, 'San Diego')

    def test_no_fake_company_from_challenge_metadata(self):
        for html in ['<title>Example Builders | ProView</title><p>Verify you are human</p>', '<h1>Access Verification</h1><p>Please enable cookies</p>']:
            self.assertTrue(access_problem(html))
            self.assertIsNone(extract_company(html, URL))
            self.assertEqual(extract_listing_companies(html, LISTING), [])

    def test_listing_details_scoped_to_each_company(self):
        cards = extract_listing_companies(CARDS, LISTING)
        self.assertEqual(len(cards), 2)
        self.assertEqual(cards[0].company_name, 'Example Builders LLC')
        self.assertEqual(cards[0].phones[0]['phone'], '+16195550110')
        self.assertEqual(cards[0].emails, ['office@example.org'])
        self.assertEqual(cards[0].city, 'San Diego')
        self.assertEqual(cards[1].phones[0]['phone'], '+15125550188')
        self.assertEqual(cards[1].emails, [])
        self.assertEqual(cards[1].state, 'TX')
        self.assertEqual(cards[0].contacts, [])

    def test_listing_records_survive_all_profile_pages_blocked(self):
        from openpyxl import load_workbook
        class BlockedWorker(FixtureWorker):
            def load(self, url, delay):
                if url != LISTING:
                    raise PermissionError('Access verification')
                return super().load(url, delay)
            def wait_manual(self, reason):
                self.stop_event.set()
                raise Stopped()
        with tempfile.TemporaryDirectory() as tmp:
            worker = BlockedWorker(queue.Queue(), tmp)
            worker.page = FixturePage({LISTING: CARDS})
            with self.assertRaises(Stopped):
                worker.crawl({'seeds': [LISTING], 'broad': False, 'enrich': False, 'retry': True, 'delay': 2, 'limit': 0})
            self.assertEqual(worker.store.stats()['companies'], 2)
            book = load_workbook(worker.store.last_excel_path)
            self.assertEqual(book['Contacts'].max_row, 3)
            self.assertIn('listing_card', [row[19].value for row in list(book['Contacts'])[1:]])
            book.close()
            worker.store.close()

    def test_zero_export_shows_summary_and_all_queued_urls(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.enqueue(URL, 'contacts')
            store.db.commit()
            self.assertEqual(store.export(), 0)
            book = load_workbook(store.last_excel_path)
            self.assertEqual(book.active.title, 'Summary')
            self.assertIn('NO CONTACTS COLLECTED', book['Summary']['B5'].value)
            self.assertEqual(book['Pages']['A2'].value, URL)
            self.assertEqual(book['Pages']['B2'].value, 'pending')
            self.assertEqual(book['Contacts'].max_row, 1)
            book.close()
            store.close()

    def test_retry_refreshes_completed_listing_and_reprocesses_failed_profiles(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.enqueue(LISTING, 'listing')
            store.enqueue(URL, 'contacts')
            store.finish(LISTING, 'done')
            store.finish(URL, 'error', 'No company heading found')
            store.retry_errors()
            self.assertEqual(store.stats()['pending'], 2)
            store.close()

    def test_cached_layout_failure_can_be_reextracted_offline(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.snapshot(URL, HTML.replace('<h1>', '<div class="companyName">').replace('</h1>', '</div>'))
            store.close()
            store = Store(tmp)
            self.assertEqual(store.recover_cached_pages(), 1)
            self.assertEqual(store.export(), 1)
            store.close()

    def test_locked_csv_does_not_prevent_current_excel_export(self):
        import os
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.save_company(extract_company(HTML, URL).dictionary())
            real_replace = os.replace
            def locked_csv(source, target):
                if Path(target).name == 'contacts.csv':
                    raise PermissionError('CSV is open')
                return real_replace(source, target)
            with patch('storage.os.replace', side_effect=locked_csv):
                self.assertEqual(store.export(), 1)
            self.assertTrue(store.last_csv_path.is_file())
            self.assertTrue(store.last_csv_path.name.startswith('contacts_saved_'))
            self.assertTrue(store.last_excel_path.is_file())
            store.close()

    def test_export_request_while_running_is_honored_at_checkpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            worker = Worker(queue.Queue(), tmp)
            worker.store = Store(tmp)
            worker.store.save_company(extract_company(HTML, URL).dictionary())
            worker.export_event.set()
            worker.check()
            self.assertFalse(worker.export_event.is_set())
            self.assertTrue(worker.store.last_excel_path.is_file())
            self.assertEqual(json.loads((Path(tmp) / 'run_summary.json').read_text())['rows'], 1)
            worker.store.close()

    def test_listing_cannot_downgrade_profile_details(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            company = extract_company(HTML, URL)
            store.save_company(company.dictionary())
            store.save_company(extract_listing_companies(CARDS, LISTING)[0].dictionary())
            saved = store.get_company(company.company_id)
            self.assertEqual(saved['company_name'], 'Example Builders LLC')
            self.assertEqual(saved['method'], 'profile_page')
            self.assertEqual(saved['source_url'], URL)
            self.assertEqual(saved['contacts'][0]['name'], 'Alex Morgan')
            store.close()


class BrowserFailureTests(unittest.TestCase):
    class Response:
        def __init__(self, status, headers=None):
            self.status, self.headers = status, headers or {}

    class Page:
        def __init__(self, responses, contents, url='about:blank'):
            self.responses, self.contents, self.index = responses, contents, -1
            self.url = url
            self.navigations = 0
        def goto(self, url, **kwargs):
            self.url = url
            self.navigations += 1
            self.index += 1
            response = self.responses[self.index]
            if isinstance(response, Exception):
                raise response
            return response
        def content(self):
            return self.contents[max(0, self.index)]
        def wait_for_timeout(self, ms):
            pass

    def worker(self, page):
        worker = Worker(queue.Queue(), '.')
        worker.page = page
        worker.pace = lambda delay: None
        worker.save_exports = lambda final=False: None
        return worker

    def test_rate_limit_cools_down_then_retries_once(self):
        page = self.Page([self.Response(429, {'retry-after': '40'}), self.Response(200)], ['<h1>Too Many Requests</h1>', HTML])
        worker = self.worker(page)
        with patch.object(worker.stop_event, 'wait', return_value=False) as wait:
            self.assertEqual(worker.load(URL, 5), HTML)
        wait.assert_called_once_with(40)
        self.assertEqual(page.navigations, 2)

    def test_persistent_block_needs_normal_browser_access(self):
        blocked = '<title>Access Verification</title><p>Verify you are human</p>'
        page = self.Page([self.Response(403)], [blocked])
        worker = self.worker(page)
        with patch.object(worker, 'wait_manual') as manual:
            with self.assertRaises(PermissionError):
                worker.load(URL, 5)
        manual.assert_called_once()
        self.assertEqual(page.navigations, 1)

    def test_manual_access_uses_page_content_without_old_403_status(self):
        page = self.Page([self.Response(403)], ['<h1>Access Verification</h1>'])
        worker = self.worker(page)
        with patch.object(worker, 'wait_manual', side_effect=lambda reason: page.contents.__setitem__(0, HTML)):
            self.assertEqual(worker.load(URL, 5), HTML)

    def test_stop_interrupts_rate_limit_cooldown(self):
        page = self.Page([self.Response(429)], ['<h1>Too many requests</h1>'])
        worker = self.worker(page)
        with patch.object(worker.stop_event, 'wait', return_value=True):
            with self.assertRaises(Stopped):
                worker.load(URL, 5)

    def test_navigation_timeout_keeps_usable_document(self):
        page = self.Page([TimeoutError('Navigation timeout')], [HTML])
        worker = self.worker(page)
        self.assertEqual(worker.load(URL, 5), HTML)
        self.assertEqual(page.navigations, 1)

    def test_browser_page_already_open_is_reused(self):
        page = self.Page([], [HTML], url=URL)
        worker = self.worker(page)
        self.assertEqual(worker.load(URL, 5), HTML)
        self.assertEqual(page.navigations, 0)


if __name__ == '__main__':
    unittest.main()
