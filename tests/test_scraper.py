import csv
import json
import sys
import tempfile
import unittest
import queue
from unittest.mock import patch
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extractor import extract_company, discover, canonical, phones, emails
from storage import Store, job_folder
from crawler import public_web_url, Worker, Stopped

URL = "https://www.thebluebook.com/iProView/123456/example-builders/subcontractors/locations-contacts/"
HTML = '''<html><body>
<header>888-720-1710 support@thebluebook.com</header>
<h1>Example Builders LLC</h1><p>San Diego, CA 92101</p>
<a href="https://example.org/">Website</a>
<h2>Locations</h2><h3>Example Builders LLC</h3>
<address itemprop="streetAddress">100 Example Road</address>
<p>County: San Diego</p>
<p>View Phone (619) 555-0110</p><p>View Fax (619) 555-0111</p>
<a href="mailto:office@example.org">Email</a>
<h2>Key Contacts</h2>
<h4>Owners, Principals &amp; Senior Executives</h4>
<h5>Alex Morgan</h5><p>Owner</p><p>View Phone</p><p>(619) 555-0120 ext. 7</p>
<a href="mailto:alex@example.org">Email</a>
<h5>Alex Morgan</h5><p>View Phone (619) 555-0120 ext. 7</p>
<h4>Estimating, Purchasing &amp; Preconstruction</h4>
<h5>Sam Rivera</h5><p>Estimator</p><p>View Phone (858) 555-0140</p>
<p>View Fax</p><p>(858) 555-0141</p>
<h5>Unknown User</h5><p>View Phone (858) 555-0199</p>
<nav>Find Companies 888-720-1710</nav>
<footer>Boston, MA 02109 (617) 555-0100</footer>
<form>Fake Form Person <input value="not_a_contact@example.com"></form>
</body></html>'''


class ExtractionTests(unittest.TestCase):
    def test_names_phones_emails_remain_associated(self):
        c = extract_company(HTML, URL)
        self.assertEqual(c.company_name, "Example Builders LLC")
        self.assertEqual((c.city, c.state, c.county), ("San Diego", "CA", "San Diego"))
        self.assertEqual(c.address, "100 Example Road")
        self.assertEqual(c.emails, ["office@example.org"])
        self.assertEqual(c.website, "https://example.org/")
        self.assertEqual(len(c.contacts), 3)
        self.assertEqual(c.contacts[0].name, "Alex Morgan")
        self.assertEqual(c.contacts[0].emails, ["alex@example.org"])
        self.assertEqual(c.contacts[0].phones, [{"phone": "+16195550120", "extension": "7", "kind": "Phone"}])
        self.assertEqual(c.contacts[1].title, "Estimator")
        self.assertEqual(c.contacts[1].phones[1]["kind"], "Fax")
        self.assertNotIn("+18585550199", [p["phone"] for p in c.contacts[1].phones])
        self.assertEqual(c.contacts[2].name, "")
        self.assertEqual(c.contacts[2].phones[0]["phone"], "+18585550199")
        self.assertNotIn("+16175550100", [p["phone"] for p in c.phones])

    def test_vendor_numbers_and_email_assets_excluded(self):
        self.assertEqual(phones("Call 888-720-1710"), [])
        self.assertEqual(emails("info@example.org logo@2x.png"), ["info@example.org"])

    def test_unpublished_is_blank_not_guessed(self):
        c = extract_company("<h1>Example Electric</h1><h2>Key Contacts</h2><h5>Jordan Smith</h5><p>CEO</p>", URL)
        self.assertEqual(c.contacts[0].phones, [])
        self.assertEqual(c.contacts[0].emails, [])

    def test_gated_page_is_not_a_company(self):
        self.assertIsNone(extract_company("<h1>Access Verification</h1><p>Proceed</p>", URL))

    def test_schema_person(self):
        structured = '<h1>Example Company</h1><script type="application/ld+json">' + json.dumps({"@graph": [{"@type": "Person", "name": "Taylor Chen", "jobTitle": "CEO", "telephone": "619-555-0150", "email": "taylor@example.org"}]}) + '</script>'
        c = extract_company(structured, URL)
        self.assertEqual(c.contacts[0].name, "Taylor Chen")
        self.assertEqual(c.contacts[0].method, "structured_person")

    def test_pagination_preserves_filters_and_nav(self):
        url = "https://www.thebluebook.com/search.html?region=3&class=2200&page=1"
        markup = '<a href="/iProView/123456/example-builders/subcontractors/">Profile</a><nav class="pagination"><a href="?region=3&class=2200&page=2">2</a><a rel="next" href="?region=3&class=2200&page=3">Next</a></nav><a href="https://evil.example.org">Elsewhere</a>'
        links = discover(markup, url)
        self.assertEqual(len(links), 3)
        self.assertIn(("https://www.thebluebook.com/search.html?class=2200&page=2&region=3", "listing"), links)

    def test_no_invented_pagination(self):
        self.assertEqual(discover("<p>83 pages</p>", "https://www.thebluebook.com/search.html?class=2200"), [])

    def test_canonical_and_public_urls(self):
        self.assertEqual(canonical("https://thebluebook.com/search.html?b=2&a=1&utm_source=x#foo"), "https://www.thebluebook.com/search.html?a=1&b=2")
        self.assertFalse(public_web_url("http://127.0.0.1/"))
        self.assertFalse(public_web_url("http://192.168.1.1/"))
        self.assertFalse(public_web_url("file:///etc/passwd"))
        self.assertFalse(public_web_url("http://localhost/"))
        self.assertTrue(public_web_url("https://example.org/"))


class PersistenceTests(unittest.TestCase):
    def test_resume_deduplication_and_atomic_exports(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.enqueue(URL, "contacts")
            store.enqueue(URL, "contacts")
            self.assertEqual(store.stats()["pending"], 1)
            store.pop()
            store.save_company(extract_company(HTML, URL).dictionary())
            store.save_company(extract_company(HTML, URL).dictionary())
            rows = list(store.rows())
            self.assertEqual(len(rows), 1)  # One preferred owner number per company.
            self.assertEqual(sum(r["Contact Name"] == "Alex Morgan" for r in rows), 1)
            self.assertEqual(next(r for r in rows if r["Contact Name"] == "Alex Morgan")["Email"], "alex@example.org")
            self.assertEqual(store.export(), 1)
            store.close()
            reopened = Store(tmp)
            self.assertEqual(reopened.stats()["pending"], 1)
            reopened.finish(URL, "blocked", "Access check")
            reopened.retry_errors()
            self.assertEqual(reopened.stats()["pending"], 1)
            reopened.close()
            numbers = (Path(tmp) / "phones.txt").read_text().splitlines()
            self.assertEqual(numbers, ["+16195550120"])
            from openpyxl import load_workbook
            book = load_workbook(Path(tmp) / "contacts.xlsx")
            self.assertEqual(book["Contacts"].max_row, 2)
            for row in book["Contacts"].iter_rows(min_row=2):
                self.assertEqual(row[4].data_type, "s")

    def test_open_excel_workbook_gets_a_readable_alternate_export(self):
        import os
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            store.export()
            store.save_company(extract_company(HTML, URL).dictionary())
            replace = os.replace
            def locked_workbook(source, target):
                if Path(target).name == "contacts.xlsx":
                    raise PermissionError("Workbook is open in Excel")
                return replace(source, target)
            with patch("storage.os.replace", side_effect=locked_workbook):
                self.assertEqual(store.export(), 1)
            self.assertTrue(store.last_excel_path.name.startswith("contacts_saved_"))
            book = load_workbook(store.last_excel_path)
            self.assertEqual(book["Contacts"].max_row, 2)
            book.close()
            store.close()

    def test_separate_jobs_and_same_job_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            a = job_folder(tmp, [URL], False, True)
            b = job_folder(tmp, [URL + "?x=2"], False, True)
            self.assertNotEqual(a, b)
            self.assertEqual(a, job_folder(tmp, [URL], False, True))


class EmptyLocator:
    def filter(self, **kwargs):
        return self

    def count(self):
        return 0

    def all(self):
        return []


class FixturePage:
    def __init__(self, pages):
        self.pages = pages
        self.url = next(iter(pages))

    def content(self):
        return self.pages[self.url]

    def evaluate(self, *args):
        return None

    def wait_for_timeout(self, *args):
        return None

    def locator(self, *args):
        return EmptyLocator()


class FixtureWorker(Worker):
    def load(self, url, delay):
        self.check()
        self.page.url = url
        return self.page.content()

    def expand_public_details(self):
        pass


class CrawlTests(unittest.TestCase):
    def test_three_page_crawl_and_completed_job_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            pages = {}
            for index in range(1, 4):
                listing = canonical(f"https://www.thebluebook.com/search.html?region=3&page={index}")
                profile = f"https://www.thebluebook.com/iProView/{123450 + index}/example-{index}/subcontractors/"
                pages[listing] = f'<a href="{profile}">Company {index}</a>' + (f'<nav class="pagination"><a rel="next" href="?region=3&page={index+1}">Next</a></nav>' if index < 3 else '')
                pages[profile] = HTML.replace("Example Builders LLC", f"Company {index} LLC")
                pages[profile + "locations-contacts/"] = pages[profile]
            worker = FixtureWorker(queue.Queue(), tmp)
            worker.page = FixturePage(pages)
            options = {"seeds": [canonical("https://www.thebluebook.com/search.html?region=3&page=1")], "broad": False, "enrich": False, "retry": True, "delay": 2, "limit": 0}
            worker.crawl(options)
            self.assertEqual(worker.store.stats()["done"], 9)
            self.assertEqual(worker.store.stats()["companies"], 3)
            folder = worker.store.folder
            worker.store.close()
            worker.store = None
            worker.crawl(options)
            self.assertEqual(worker.store.stats()["done"], 9)
            self.assertEqual(worker.store.stats()["companies"], 3)
            self.assertEqual(worker.store.folder, folder)
            worker.store.close()

    def test_stop_at_500_of_1000_exports_partial_data_and_can_resume(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            pages = {}
            for index in range(1000):
                url = f"https://www.thebluebook.com/iProView/{100000 + index}/example/subcontractors/locations-contacts/"
                pages[url] = f"<h1>Company {index}</h1><p>Phone (619) 555-0110</p>"
            class StopWorker(FixtureWorker):
                loaded = 0
                def load(self, url, delay):
                    if self.loaded == 500:
                        self.stop_event.set()
                    self.check()
                    self.loaded += 1
                    return super().load(url, delay)
            worker = StopWorker(queue.Queue(), tmp)
            worker.page = FixturePage(pages)
            options = {"seeds": list(pages), "broad": False, "enrich": False, "retry": False, "delay": 2, "limit": 0, "autosave": False}
            # Keep this large interruption test focused on the final export;
            # periodic exports are checked separately below.
            save = worker.save_exports
            worker.save_exports = lambda final=False: save(final=True) if final else None
            with self.assertRaises(Stopped):
                worker.crawl(options)
            self.assertEqual(worker.store.stats()["done"], 500)
            self.assertEqual(worker.store.stats()["companies"], 500)
            folder = worker.store.folder
            book = load_workbook(folder / "contacts.xlsx")
            self.assertEqual(book["Contacts"].max_row, 2)  # All fixture companies share one number.
            self.assertIn("Company 499", book["Contacts"]["A2"].value)
            book.close()
            with (folder / "contacts.csv").open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 1)
            self.assertIn("Stopped", next(value for kind, value in list(worker.events.queue) if kind == "finished"))
            worker.store.close()
            resumed = Store(folder)
            self.assertEqual(resumed.stats()["done"], 500)
            self.assertEqual(resumed.stats()["companies"], 500)
            self.assertGreater(resumed.stats()["pending"], 0)
            resumed.close()

    def test_every_page_option_exports_before_job_finishes(self):
        with tempfile.TemporaryDirectory() as tmp:
            worker = FixtureWorker(queue.Queue(), tmp)
            worker.page = FixturePage({URL: HTML})
            options = {"seeds": [URL], "broad": False, "enrich": False, "retry": True, "delay": 2, "limit": 1, "autosave": True}
            with patch.object(worker, "save_exports", wraps=worker.save_exports) as save:
                worker.crawl(options)
                self.assertEqual(save.call_count, 2)  # one page plus final save
            self.assertTrue((worker.store.folder / "contacts.xlsx").exists())
            worker.store.close()

    def test_stop_during_website_check_preserves_collected_emails(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            profile = URL.removesuffix("locations-contacts/")
            class EnrichmentWorker(FixtureWorker):
                def enrich(self, data, delay):
                    data["emails"].append("new@example.org")
                    data["email_sources"].append("https://example.org/contact")
                    self.stop_event.set()
                    self.check()
            worker = EnrichmentWorker(queue.Queue(), tmp)
            worker.page = FixturePage({profile: HTML, profile + "locations-contacts/": HTML})
            with self.assertRaises(Stopped):
                worker.crawl({"seeds": [profile], "broad": False, "enrich": True, "retry": False, "delay": 2, "limit": 0})
            book = load_workbook(worker.store.folder / "contacts.xlsx")
            header = [cell.value for cell in book["Contacts"][1]]
            email_column = header.index("Company Email")
            self.assertIn("new@example.org", list(book["Contacts"].values)[1][email_column])
            book.close()
            worker.store.close()

    def test_stop_while_paused_still_exports_previous_contacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            worker = FixtureWorker(queue.Queue(), tmp)
            worker.store = Store(tmp)
            worker.store.save_company(extract_company(HTML, URL).dictionary())
            worker.pause_event.set()
            worker.stop_event.set()
            with patch.object(worker, "_crawl", side_effect=lambda options: worker.check()):
                with self.assertRaises(Stopped):
                    worker.crawl({})
            self.assertTrue((Path(tmp) / "contacts.xlsx").exists())
            worker.store.close()

    def test_session_limit_preserves_remaining_queue(self):
        with tempfile.TemporaryDirectory() as tmp:
            listing = "https://www.thebluebook.com/search.html"
            profile = "https://www.thebluebook.com/iProView/123456/example-builders/subcontractors/"
            worker = FixtureWorker(queue.Queue(), tmp)
            worker.page = FixturePage({listing: f'<a href="{profile}">Company</a>'})
            worker.crawl({"seeds": [listing], "broad": False, "enrich": False, "retry": True, "delay": 2, "limit": 1})
            self.assertEqual(worker.store.stats()["done"], 1)
            self.assertEqual(worker.store.stats()["pending"], 2)
            worker.store.close()


if __name__ == "__main__":
    unittest.main()
